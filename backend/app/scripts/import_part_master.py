r"""
부품마스터 임포트 스크립트.

Engine_parts_data_filled_v2.xlsx 의 '치수Data' 시트(원자재 치수/무게)와
'Sheet2_포장박스예측' 시트(개별포장 분류/치수/완충재)를 품번 기준으로
조인해서 PartType 테이블에 일괄 등록한다.

실행 (프로젝트 루트에서):
    python -m app.scripts.import_part_master "C:\경로\Engine_parts_data_filled_v2.xlsx"
"""

import re
import sys
import time
from collections import defaultdict

import openpyxl
from sqlalchemy.orm import Session

from app.database import Base, engine, SessionLocal
from app.models import PartType
from app.scripts.size_parser import (
    parse_size, match_category, is_cube_placeholder, rescale_to_ratio,
)


def _normalize_part_no(part_no: str) -> str:
    """품번 표기(공백/점/하이픈 혼용) 정규화 - 중복 검출/매칭용"""
    return str(part_no).strip().replace(" ", "").replace(".", "").replace("-", "").upper()


def load_valid_service_parts(xlsx_path: str) -> set:
    """[2026-10 추가] EG_PARTSEXCLUDE_NOT_SERVICE.XLSX ('Material'/'Service Det.' 컬럼) -
    현재 실제 사용 중인 서비스품목 품번 목록. 이 목록에 없는 품번은 마스터에서 제외한다
    (총 53,271건 중 25,935건만 이 목록에 있고, 나머지 27,336건은 미사용/구품번으로 확인됨)."""
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb.active
    valid = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        valid.add(_normalize_part_no(row[0]))
    return valid


_REVISION_SUFFIX_PATTERN = re.compile(r"^(.*\d)([A-Z])$")


def find_safe_duplicate_removals(ws1) -> set:
    """
    품번이 '숫자...+알파벳 1글자'로 끝나고, 그 알파벳만 뗀 버전(base)도
    목록에 있는 경우를 찾는다. 이때 품명/무게/치수가 전부 완전히 똑같은
    경우에만 '같은 부품의 리비전 표기 차이'로 보고, 알파벳 없는(구버전)
    쪽을 제거 대상으로 돌려준다.
    """
    groups = defaultdict(list)
    for row in ws1.iter_rows(min_row=2, values_only=True):
        part_no, name, _suffix, _model, _desc2, netw, _grossw, size_text, _confidence, _evidence = row
        if not part_no:
            continue
        pn = str(part_no).strip()
        m = _REVISION_SUFFIX_PATTERN.match(pn)
        base, letter = (m.group(1), m.group(2)) if m else (pn, "")
        groups[base].append({"letter": letter, "pn": pn, "name": name, "netw": netw, "size": size_text})

    to_remove = set()
    for base, variants in groups.items():
        if len(variants) < 2:
            continue
        names = {v["name"] for v in variants}
        weights = {v["netw"] for v in variants}
        sizes = {v["size"] for v in variants}
        if len(names) != 1 or len(weights) != 1 or len(sizes) != 1:
            continue
        lettered = [v for v in variants if v["letter"]]
        if not lettered:
            continue
        keep = max(lettered, key=lambda v: v["letter"])
        for v in variants:
            if v["pn"] != keep["pn"]:
                to_remove.add(v["pn"])
    return to_remove


def build_category_ratios(ws1, min_samples: int = 5) -> dict:
    """'신뢰도 상'(실측) 데이터만 모아서 카테고리별 평균 가로:세로:높이 비율을 계산한다."""
    samples = defaultdict(list)
    for row in ws1.iter_rows(min_row=2, values_only=True):
        part_no, name, _suffix, _model, _desc2, _netw, _grossw, size_text, confidence, _evidence = row
        if confidence != "상":
            continue
        dims = parse_size(size_text)
        if not dims:
            continue
        kw = match_category(name)
        if not kw:
            continue
        l, w, h = sorted(dims, reverse=True)
        if h <= 0:
            continue
        samples[kw].append((l / h, w / h))

    ratios = {}
    for kw, vals in samples.items():
        if len(vals) < min_samples:
            continue
        avg_l = sum(v[0] for v in vals) / len(vals)
        avg_w = sum(v[1] for v in vals) / len(vals)
        ratios[kw] = (avg_l, avg_w, 1.0)
    return ratios


def _improve_if_placeholder(dims, confidence, name, category_ratios):
    """'신뢰도 하' + 정육면체 가정치로 보이는 치수를 같은 카테고리 실측비율로 재배분"""
    if confidence != "하" or not is_cube_placeholder(dims):
        return dims, False
    kw = match_category(name)
    if not kw or kw not in category_ratios:
        return dims, False
    return rescale_to_ratio(dims, category_ratios[kw]), True


def load_packaging_map(ws2, category_ratios: dict) -> dict:
    """Sheet2_포장박스예측 시트를 품번 기준 dict로 로드 (정육면체 추정치는 비율 보정)"""
    packaging_map = {}
    for row in ws2.iter_rows(min_row=2, values_only=True):
        part_no, name, category, size_text, cushioning, confidence, _evidence = row
        if not part_no:
            continue
        key = _normalize_part_no(part_no)
        size = parse_size(size_text)
        improved_size, improved = (size, False)
        if size:
            improved_size, improved = _improve_if_placeholder(size, confidence, name, category_ratios)
        packaging_map[key] = {
            "category": category,
            "size": improved_size,
            "cushioning": cushioning,
            "dimension_improved": improved,
        }
    return packaging_map


def import_part_master(xlsx_path: str, db: Session, batch_size: int = 1000,
                        valid_parts: set | None = None) -> dict:
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws1 = wb["치수Data"]
    ws2 = wb["Sheet2_포장박스예측"]

    print("카테고리별 실측(신뢰도 '상') 비율 계산 중...")
    category_ratios = build_category_ratios(ws1)
    print(f"  -> {len(category_ratios)}개 카테고리 비율 확보: {sorted(category_ratios.keys())}")

    wb_dup = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    print("끝알파벳만 다른 안전한 중복 품번 탐지 중 (값이 완전히 동일한 경우만)...")
    duplicates_to_skip = find_safe_duplicate_removals(wb_dup["치수Data"])
    print(f"  -> {len(duplicates_to_skip)}건 제거 대상 (구버전 리비전 표기)")

    wb2 = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws1 = wb2["치수Data"]
    ws2 = wb2["Sheet2_포장박스예측"]

    packaging_map = load_packaging_map(ws2, category_ratios)

    stats = {
        "total": 0, "created": 0, "updated": 0, "unparsed_size": 0,
        "dimension_improved": 0, "duplicate_skipped": 0, "excluded_not_service": 0,
    }
    existing_ids = {p.part_no: p for p in db.query(PartType).all()}

    buffer = []
    started = time.time()

    for row in ws1.iter_rows(min_row=2, values_only=True):
        part_no, name, _suffix, engine_model, _desc2, net_w, gross_w, size_text, confidence, evidence = row
        if not part_no:
            continue
        stats["total"] += 1

        pn_str = str(part_no).strip()
        if pn_str in duplicates_to_skip:
            stats["duplicate_skipped"] += 1
            continue

        if valid_parts is not None and _normalize_part_no(pn_str) not in valid_parts:
            stats["excluded_not_service"] += 1
            continue

        size = parse_size(size_text)
        if size is None:
            stats["unparsed_size"] += 1
        else:
            size, improved = _improve_if_placeholder(size, confidence, name, category_ratios)
            if improved:
                stats["dimension_improved"] += 1
                evidence = (evidence or "") + f" | [치수 재추정: {match_category(name)} 카테고리 실측비율 적용, 정육면체 가정치 대체]"
        length, width, height = size if size else (None, None, None)

        pack_info = packaging_map.get(_normalize_part_no(part_no), {})
        pack_size = pack_info.get("size")
        pkg_length, pkg_width, pkg_height = pack_size if pack_size else (None, None, None)

        existing = existing_ids.get(part_no)
        if existing:
            existing.name = name
            existing.engine_model = engine_model
            existing.net_weight_kg = net_w
            existing.gross_weight_kg = gross_w
            existing.raw_length_mm = length
            existing.raw_width_mm = width
            existing.raw_height_mm = height
            existing.confidence_level = confidence
            existing.evidence = evidence
            existing.packaging_category = pack_info.get("category")
            existing.packaging_length_mm = pkg_length
            existing.packaging_width_mm = pkg_width
            existing.packaging_height_mm = pkg_height
            existing.cushioning_material = pack_info.get("cushioning")
            stats["updated"] += 1
        else:
            buffer.append(
                PartType(
                    part_no=part_no,
                    name=name,
                    engine_model=engine_model,
                    net_weight_kg=net_w,
                    gross_weight_kg=gross_w,
                    raw_length_mm=length,
                    raw_width_mm=width,
                    raw_height_mm=height,
                    confidence_level=confidence,
                    evidence=evidence,
                    packaging_category=pack_info.get("category"),
                    packaging_length_mm=pkg_length,
                    packaging_width_mm=pkg_width,
                    packaging_height_mm=pkg_height,
                    cushioning_material=pack_info.get("cushioning"),
                )
            )
            stats["created"] += 1

        if len(buffer) >= batch_size:
            db.bulk_save_objects(buffer)
            db.commit()
            buffer = []
            elapsed = time.time() - started
            print(f"  ... {stats['total']}건 처리 ({elapsed:.1f}초 경과)")

    if buffer:
        db.bulk_save_objects(buffer)
        db.commit()

    return stats


def main():
    if len(sys.argv) < 2:
        print("사용법: python -m app.scripts.import_part_master <엑셀파일경로> [서비스품목목록.xlsx]")
        print("  두 번째 인자(EG_PARTSEXCLUDE_NOT_SERVICE.XLSX 등)를 주면,")
        print("  그 목록(Material 컬럼)에 없는 품번은 마스터에서 제외합니다.")
        sys.exit(1)

    xlsx_path = sys.argv[1]
    valid_parts = None
    if len(sys.argv) >= 3:
        valid_list_path = sys.argv[2]
        print(f"서비스품목 목록 로딩: {valid_list_path}")
        valid_parts = load_valid_service_parts(valid_list_path)
        print(f"  -> {len(valid_parts)}건 (이 목록에 없는 품번은 마스터에서 제외됩니다)")

    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        print(f"부품마스터 임포트 시작: {xlsx_path}")
        stats = import_part_master(xlsx_path, db, valid_parts=valid_parts)
        print()
        print("=== 임포트 완료 ===")
        print(f"총 처리:        {stats['total']}건")
        print(f"신규 생성:      {stats['created']}건")
        print(f"기존 갱신:      {stats['updated']}건")
        print(f"치수 파싱실패:  {stats['unparsed_size']}건 (수작업 검토 대상, raw_*_mm이 비어있음)")
        print(f"치수 재추정:    {stats['dimension_improved']}건 (정육면체 가정치 -> 카테고리 실측비율로 보정)")
        print(f"중복품번 제거:  {stats['duplicate_skipped']}건 (끝알파벳만 다르고 값 완전동일한 구버전 리비전)")
        print(f"서비스품목 아님(제외): {stats['excluded_not_service']}건")
    finally:
        db.close()


if __name__ == "__main__":
    main()
